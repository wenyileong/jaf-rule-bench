"""Acceptance tests (specification s 9; build plan ss 9, 10, 12, 13 and 14).
Run: PYTHONPATH=. python3 tests/test_acceptance.py
"""
import numpy as np
from jaf.params import Params, RULES
from jaf.rules import Compiled
from jaf.equilibrium import solve
from jaf.network import resolvent, round_trip
from jaf import experiments as X

results = []
def check(name, cond, detail=""):
    results.append(bool(cond)); print(("PASS " if cond else "FAIL ") + name + (f"  [{detail}]" if detail else ""))

rng = np.random.default_rng(0)
F = rng.random((5, 5)); np.fill_diagonal(F, 0); D = rng.random(5) * 3
Pm, M = resolvent(F, D, 0.6); S = np.eye(5); T = np.eye(5)
for _ in range(400): T = T @ (0.6 * Pm); S = S + T
check("1 resolvent = power series", np.abs(S - M).max() < 1e-12)

full = dict(obs_custodial=1.0, obs_self=1.0, sigma=0.0)
c = Compiled("owner_residence", Params(dual_residence=0.0, **full)); e = solve(c, "a", tol=1e-8)
Sj = np.bincount(c.j, weights=e["s"], minlength=6)
err = np.abs(e["tau"] - 1 / (e["theta"] * (1 - Sj))).max()
check("2 partition reproduces tau = 1/[theta(1 - s)] (plain logit)", err < 1e-4, f"{err:.1e}")

c = Compiled("enclosure", Params()); sols = [solve(c, "a", tau0=rng.uniform(0.1, 0.9, 6))["tau"] for _ in range(4)]
spread = np.ptp(np.array(sols), axis=0).max()
check("3 equilibrium unique across starts", spread < 1e-3, f"{spread:.1e}")

P = Params()
res = {r: dict(zip(("comp", "a", "b"), X.run_rule(r, P))) for r in RULES}
check("4 all rules converge under both concepts", all(res[r]["a"]["converged"] and res[r]["b"]["converged"] for r in RULES))
check("5 only non-assignment is a corner", [r for r in RULES if res[r]["a"]["corner"]] == ["non_assignment"])
rk = X.ranking(res); ok = True
for (r1, r2), v in rk.items():
    if v in (">", "<"):
        ok &= np.sign(res[r1]["a"]["effective"] - res[r2]["a"]["effective"]) == np.sign(res[r1]["b"]["effective"] - res[r2]["b"]["effective"])
check("6 ranking uses effective rates and requires agreement", ok)

L = {r: X.leakage(res[r]["comp"], res[r]["a"]) for r in RULES if X.usable(res[r]["a"])}
check("7 leakage decomposition sums to one", all(abs(sum(l[k] for k in X.LEAK_KEYS) - 1) < 1e-6 for l in L.values()))
o = L["owner_residence"]
check("8 owner-residence: no round-tripping, conduit switching only via dual residents (<1%), exit separated", o["round_trip"] < 1e-6 and o["conduit_switching"] < 0.01 and o["exit"] > 0.1,
      f"exit {o['exit']:.2f}")
check("9 intermediary location leaks into the gap", L["intermediary"]["to_gap"] > 0.2, f"{L['intermediary']['to_gap']:.2f}")

ex = solve(Compiled("X_gap", P), "a"); ey = solve(Compiled("Y_overlap", P), "a")
check("10 matched pair: X has gap, no overlap; Y the reverse", ex["failure"]["overlap"] == 0 and ey["failure"]["gap"] == 0
      and ex["failure"]["gap"] > 0 and ey["failure"]["overlap"] > 0)
z0 = round_trip(resolvent(np.ones((4, 4)) - np.eye(4), np.full(4, 10.0), 0.6)[1])[0]
F2 = np.ones((4, 4)) - np.eye(4); F2[0, 1] += 3; F2[1, 0] += 3
check("11 round trip raises Z", round_trip(resolvent(F2, np.full(4, 10.0), 0.6)[1])[0] > z0)
ei = res["intermediary"]["a"]; fi = ei["failure"]
ident = abs(ei["effective"] * ei["Q"].sum() + fi["backstop"] - ei["revenue"].sum())
check("12 effective rate excludes gap backstop; yield counts the gap", ident < 1e-9 and ei["revenue_per_base"] < ei["effective"]
      and abs(res["control"]["a"]["revenue_per_base"] - res["control"]["a"]["effective"]) < 1e-9, f"identity {ident:.1e}")

# ---- 0.3: mixed regimes (build plan s 12.2) -------------------------------------------
from jaf.mixed import MixedCompiled
from jaf.synthetic import uniform_world
from jaf.world import World, Jurisdiction, Link, Settings, illustrative_world, mismatch_matrix, claimants
from jaf.template import write_template, read_template
import io
from openpyxl import load_workbook

dev = 0.0
for r in ("owner_residence", "intermediary", "control", "enclosure", "registration"):
    w_, pw, a_ = uniform_world(r, P)
    m_ = MixedCompiled(w_, P, placement_weight=pw, holder_base=a_)
    for k in "ab":
        dev = max(dev, np.abs(solve(Compiled(r, P), k)["tau"] - solve(m_, k)["tau"]).max())
check("13 mixed engine with one common rule reproduces 0.2 equilibria", dev < 1e-6, f"max dev {dev:.1e}")

J3 = [Jurisdiction("X", "X", liability_rule="intermediary", rate_mode="fixed", effective_rate=0.3, info_exchange_from=2020, visibility_domestic_custodial=1.0, visibility_self_hosted=1.0),
      Jurisdiction("Y", "Y", liability_rule="residence", rate_mode="fixed", effective_rate=0.2, info_exchange_from=2020, visibility_domestic_custodial=1.0, visibility_self_hosted=1.0),
      Jurisdiction("Z", "Z", liability_rule="residence_citizenship", rate_mode="fixed", effective_rate=0.25, info_exchange_from=2020, visibility_domestic_custodial=1.0, visibility_self_hosted=1.0)]
w3 = World(J3, [Link("Y", "Z", "citizenship", 0.5)], Settings(scenario_year=2027, visibility_exchanged=1.0, treaty_relief=0.5))
cl = claimants(w3, 1, "custodial", 0, [(2, "citizenship")])
m3 = MixedCompiled(w3); e3 = solve(m3, "a")
s3 = e3["s"]; total = m3.revenue(e3["tau"], s3).sum()
b, pr, se, gp = m3.entry_terms(e3["tau"]); mw = s3[m3.ek] * m3.ew
check("14 three claimants: order placement, residence, citizenship; revenue adds up",
      [k for k, _ in cl] == [0, 1, 2] and abs(total - (mw * (pr + se + gp)).sum()) < 1e-12 and 1 + m3.smask.sum(1).max() == 3)

wi = illustrative_world(); mi = MixedCompiled(wi); ei = solve(mi, "b")
fx = mi.fixed
check("15 fixed-rate states keep their coded rate; competitive states converge",
      np.allclose(ei["tau"][fx], mi.fixed_tau[fx]) and ei["converged"] and not np.allclose(ei["tau"][~fx], 0.3))

mmd = mismatch_matrix(wi, "custodial"); colix = {(c["base"], c["seat"]): n for n, c in enumerate(mmd["cols"])}
mm = [[row[colix[(code, code)]] for code in wi.codes] for row in mmd["cells"]]   # co-located exchange types
iC, iA, iE, iH1 = (wi.index(c) for c in ("C", "A", "E", "H1"))
check("16 mismatch matrix matches hand-worked cells",
      mm[iC][iA]["n"] == 0 and mm[iA][iC]["claimants"] == ["C", "A"] and mm[iH1][iE]["claimants"] == ["E"] and mm[iA][iA]["claimants"] == ["A"])

b1 = write_template(kind="illustrative"); w1, iss1 = read_template(b1)
check("17 template round trip reproduces the illustrative world", w1 is not None and w1.to_dict() == wi.to_dict() and not iss1)

wb = load_workbook(io.BytesIO(b1)); sh = wb["Jurisdictions"]
H18 = [str(c.value).rstrip(" *") for c in sh[1]]; C18 = lambda name: H18.index(name) + 1   # columns by header, not position
sh.cell(4, C18("Code")).value = "A"                                   # duplicate code
sh.cell(5, C18("Liability rule")).value = "nonsense_rule"             # bad liability rule
sh.cell(6, C18("Effective rate on returns (0 to 1)")).value = None    # fixed rate missing (row 6 is D, fixed); cell(r, c, None) would not clear
sh.cell(7, C18("Verified")).value = "no"; sh.cell(7, C18("Source (OSCOLA)")).value = None   # unverified, no source
bio = io.BytesIO(); wb.save(bio); w2, iss2 = read_template(bio.getvalue())
msgs = " | ".join(i["message"] for i in iss2)
check("18 validation catches duplicates, bad values and missing rates; flags unverified and unsourced rows",
      w2 is None and "more than once" in msgs and "Liability rule" in msgs and "Effective rate" in msgs and "Not yet verified" in msgs and "No source" in msgs)

V_on = wi.custodial_visibility(); w_off = World.from_dict(wi.to_dict()); w_off.settings.scenario_year = 2000
V_off = w_off.custodial_visibility()
ia, ie = wi.index("A"), wi.index("E")
check("19 exchange of information raises visibility of foreign exchange holdings",
      V_on[ia, ie] == wi.settings.visibility_exchanged and V_off[ia, ie] == wi.jurisdictions[ia].visibility_self_hosted and V_on[ia, ia] == V_off[ia, ia])


# ---- 0.4: legal seat against operational base; treaty relief by pair (build plan s 13.2) ----
from jaf.world import Treaty
def two(test_x, test_y):
    J = [Jurisdiction("X", "X", liability_rule="intermediary", intermediary_test=test_x),
         Jurisdiction("Y", "Y", liability_rule="intermediary", intermediary_test=test_y, offshore_incorporation_in="X", offshore_incorporation_share=0.5),
         Jurisdiction("Z", "Z", liability_rule="none")]
    return World(J)
wo, wg = two("seat", "base"), two("base", "seat")
z = 2   # holder resident in Z (claims nothing itself); exchange based in Y (index 1), incorporated in X (index 0)
co = [wo.codes[k] for k, _ in claimants(wo, z, "custodial", 1, [], seat=0)]
cg = [wg.codes[k] for k, _ in claimants(wg, z, "custodial", 1, [], seat=0)]
same = [wo.codes[k] for k, _ in claimants(wo, z, "custodial", 1, [], seat=1)]
check("20 seat and base: seat-testing seat state plus base-testing base state overlap; the reverse is a gap",
      co == ["Y", "X"] and cg == [] and same == ["Y"], f"overlap {co}, gap {cg}, co-located {same}")

def relief_world(r):
    J = [Jurisdiction("A", "A", liability_rule="intermediary", intermediary_test="base", rate_mode="fixed", effective_rate=0.3,
                      info_exchange_from=2020, visibility_domestic_custodial=1.0, visibility_self_hosted=1.0),
         Jurisdiction("B", "B", liability_rule="residence", rate_mode="fixed", effective_rate=0.2,
                      info_exchange_from=2020, visibility_domestic_custodial=1.0, visibility_self_hosted=1.0)]
    return World(J, [], Settings(scenario_year=2027, visibility_exchanged=1.0, treaty_relief=0.5), [Treaty("B", "A", r, False)])
res21 = {}
for r in (1.0, 0.0):
    mw_ = MixedCompiled(relief_world(r)); ew_ = solve(mw_, "a")
    res21[r] = ew_["failure"]["unrelieved"] / max(1e-12, ew_["failure"]["overlap"])
kap = Params().kappa_relief
check("21 pair relief: full relief leaves only the smoothing residue; none leaves the full second tax",
      res21[1.0] < np.log1p(np.exp(-0.1 * kap)) / kap + 1e-9 and abs(res21[0.0] - 0.2) < 1e-9, f"full {res21[1.0]:.4f}, none {res21[0.0]:.4f}")

wb = load_workbook(io.BytesIO(write_template(kind="illustrative"))); sh = wb["Jurisdictions"]
hdr = [str(c.value).rstrip(" *") for c in sh[1]]
col = lambda name: hdr.index(name) + 1
sh.cell(3, col("Exchanges here incorporated offshore in (code)")).value = "QQ"          # unknown code
sh.cell(5, col("Share incorporated offshore (0 to 1)")).value = None                   # C: code given, share missing
tw = wb["Treaties"]; tw.cell(3, 3).value = 1.7                                          # relief out of range
bio = io.BytesIO(); wb.save(bio); w22, iss22 = read_template(bio.getvalue())
m22 = " | ".join(i["message"] for i in iss22)
check("22 validation catches unknown offshore code, missing share and relief out of range",
      w22 is None and "'QQ' is not a code" in m22 and "is needed when an offshore place" in m22 and "Relief must be between 0 and 1" in m22)


wb23 = load_workbook(io.BytesIO(write_template(kind="illustrative"))); sh23 = wb23["Jurisdictions"]
for name in ("Share incorporated offshore (0 to 1)", "Exchanges here incorporated offshore in (code)", "Intermediary located by"):
    hdr23 = [str(c.value).rstrip(" *") for c in sh23[1]]; sh23.delete_cols(hdr23.index(name) + 1)
del wb23["Treaties"]
bio = io.BytesIO(); wb23.save(bio); w23, iss23 = read_template(bio.getvalue())
check("23 a template from 0.3 (no new columns, no Treaties sheet) loads, with warnings and defaults",
      w23 is not None and not w23.treaties and all(i["level"] == "warning" for i in iss23)
      and any("Assumed to locate intermediaries by legal seat" in i["message"] for i in iss23))

# ---- version 1.0 (build plan s 14) ------------------------------------------------------------
import copy, json, math
from jaf.design import Design, Observation, ParamSpec, apply_episode, fingerprint
from jaf.measure import Measured, QUANTITIES, observation_value
from jaf.estimate import Model, identification, estimate, validate, chi2_sf
from jaf.examples import illustrative_design, illustrative_design_skeleton, TRUTH
from jaf.recovery import recovery
from jaf import prereg as PRE
from jaf.template import read_workbook

base24 = Params(); rng24 = np.random.default_rng(base24.seed * 31 + 5); lev = [0.25, 0.5, 1.0, 1.5]; P24 = []
for d_ in range(10):
    P24.append(base24.replace(seed=int(rng24.integers(1, 10_000)), mobility_scale=lev[d_ % 4], beta=float(rng24.uniform(0.0, 50.0)),
                              theta0=float(rng24.uniform(2.6, 4.0)), rho=float(rng24.uniform(0.3, 1.0)), phi=float(rng24.uniform(0.0, 0.4)),
                              sigma=float(rng24.uniform(0.5, 2.0)), kappa_relief=float(rng24.choice([10.0, 20.0, 40.0]))))
e24 = solve(Compiled("owner_residence", P24[5]), "a"); e24b = solve(Compiled("control", P24[9]), "a")
check("24 flat best responses: cycling rates accepted as an epsilon-equilibrium, and the forgone gain reported when not",
      e24["converged"] and e24["epsilon"] is not None and e24["epsilon"] < 1e-4 and e24b["epsilon"] is not None and e24b["epsilon"] < 1e-3,
      f"accepted gain {e24['epsilon']:.1e}; unconverged case forgoes {e24b['epsilon']:.1e}")

wI = illustrative_world(); dI = illustrative_design(wI)
cI = MixedCompiled(wI); eI = solve(cI, "a"); M0 = Measured(cI, eI); M1 = Measured(cI, eI, dI)
shares = sum(M0.value("custodial_share_by_base", k) for k in wI.codes)
raw = np.array([eI["s"][(cI.c == 0) & (cI.p == k)].sum() for k in range(wI.N)])
wash = np.zeros(wI.N); wash[wI.index("H1")] = 0.30; wash[wI.index("H2")] = 0.20
adj = raw / (1 - wash); want = adj[wI.index("H1")] / adj.sum()
check("25 measurement: exchange shares sum to 1; wash trading inflates reported volume by 1/(1 - wash share)",
      abs(shares - 1) < 1e-12 and abs(M1.value("custodial_share_by_base", "H1") - want) < 1e-12, f"H1 {M0.value('custodial_share_by_base', 'H1'):.4f} -> {want:.4f}")

wAll = World.from_dict(wI.to_dict())
for j_ in wAll.jurisdictions:
    j_.info_exchange_from = 2027
wAll.settings.scenario_year = 2030
cA = MixedCompiled(wAll); eA = solve(cA, "a"); MA = Measured(cA, eA)
check("26 measurement: with every state exchanging information, no resident's trading is outside the reporting net",
      max(MA.value("resident_offshore_unexchanged_share", k) for k in wAll.codes) < 1e-12)

ref = json.dumps(wI.to_dict(), sort_keys=True)
wr = apply_episode(wI, dI.episode("reform")); ww = apply_episode(wI, dI.episode("wave_2030"))
check("27 episodes change only what they state and leave the coded world untouched",
      wr.jurisdictions[0].effective_rate == 0.36 and wr.jurisdictions[0].collection == "domestic_withholding"
      and ww.settings.scenario_year == 2030 and ww.jurisdictions[wI.index("H2")].info_exchange_from == 2030
      and json.dumps(wI.to_dict(), sort_keys=True) == ref)

w28, d28, i28 = read_workbook(write_template(kind="illustrative"))
wb28 = load_workbook(io.BytesIO(write_template(kind="illustrative")))
for n_ in ("Episodes", "Data", "Measurement", "Estimation", "Hypotheses", "Quantities"):
    del wb28[n_]
bio = io.BytesIO(); wb28.save(bio); w28b, d28b, i28b = read_workbook(bio.getvalue())
check("28 design round trip through the workbook; a 0.4 workbook (no design sheets) loads with no design and no error",
      d28 is not None and d28.to_dict() == dI.to_dict() and w28b is not None and d28b is None and not [i for i in i28b if i["level"] == "error"])

wb29 = load_workbook(io.BytesIO(write_template(kind="illustrative"))); ds = wb29["Data"]
hdr29 = [str(c.value).rstrip(" *") for c in ds[1]]; col29 = lambda n: hdr29.index(n) + 1
ds.cell(3, col29("Quantity")).value = "volume_of_everything"
ds.cell(4, col29("Standard error")).value = None
ds.cell(5, col29("Episode")).value = "no_such_episode"
ds.cell(7, col29("Compared with episode")).value = None                  # T5: exit_share needs a comparison
es29 = wb29["Episodes"]; es29.cell(3, 3).value = "holder_name"           # field an episode may not change
bio = io.BytesIO(); wb29.save(bio); w29, d29, i29 = read_workbook(bio.getvalue())
m29 = " | ".join(i["message"] for i in i29)
check("29 design validation: unknown quantity, missing standard error, undefined episode, level not usable, bad episode field",
      w29 is None and "Unknown quantity" in m29 and "positive standard error" in m29 and "is not defined in the Episodes" in m29
      and "has no observed counterpart in levels" in m29 and "An episode cannot change 'holder_name'" in m29)

d30 = copy.deepcopy(dI); d30.observations = [o for o in d30.observations if o.id == "T1"]
id30a = identification(Model(wI, d30))
d30b = copy.deepcopy(dI); d30b.params = [ParamSpec("exit_weight", True, 0.6, 0.1, 2.0)]
id30b = identification(Model(wI, d30b))
id30c = identification(Model(wI, dI))
check("30 identification: order condition, an irrelevant parameter and a well-identified design are told apart",
      id30a["status"] == "fail" and "order condition" in " ".join(id30a["reasons"])
      and id30b["status"] == "fail" and id30b["params"][0]["status"] in ("not identified", "no effect") and id30c["status"] == "ok",
      f"{id30a['status']} / {id30b['params'][0]['status']} / {id30c['status']}")

r31 = estimate(Model(wI, d30b), check_first=True)
check("31 estimation refuses when identification fails, unless overridden", r31["status"] == "refused")

rec = recovery(wI, illustrative_design_skeleton(), TRUTH, reps=1, seed=5)
s32 = {s_["name"]: s_ for s_ in rec["summary"]}
check("32 known-truth recovery: noiseless estimates within 0.1% of range; noisy estimate within 3.5 s.e.",
      rec["passed"] and all(abs(s_["noiseless_rel_error"]) < 1e-3 for s_ in rec["summary"]),
      "; ".join(f"{n}: {s_['noiseless']:.4f} (truth {s_['truth']}), noisy z {s_['max_abs_z']:.2f}" for n, s_ in s32.items()))
check("33 validation: with noiseless data, every holdout prediction is consistent", rec["runs"][0]["validation"]["passed"] == rec["runs"][0]["validation"]["n"] == 3)

r34 = PRE.freeze(wI, illustrative_design_skeleton(), author="Test")
rt = json.loads(json.dumps(r34))
c34a = PRE.check(rt, wI, dI)                                             # values added later: still consistent
d34 = copy.deepcopy(dI); d34.observations[0].sd = 0.05; d34.params[0].upper = 7.0
c34b = PRE.check(rt, wI, d34)
rt2 = copy.deepcopy(rt); rt2["content"]["observations"][0]["sd"] = 0.5
c34c = PRE.check(rt2, wI, dI)
md = PRE.to_markdown(rt)
check("34 pre-registration: values excluded from the fingerprint; changes after freezing and edits to the record detected",
      c34a["consistent"] and not r34["values_present_at_freeze"] and len(c34b["deviations"]) == 2 and not c34c["record_intact"]
      and r34["fingerprint"] in md and fingerprint({k: v for k, v in rt.items() if k != "fingerprint"}) == r34["fingerprint"])

check("35 chi-square tail probabilities", abs(chi2_sf(3.841459, 1) - 0.05) < 1e-5 and abs(chi2_sf(11.0705, 5) - 0.05) < 1e-5
      and abs(chi2_sf(2.0, 2) - math.exp(-1)) < 1e-10)

# ---- independent review findings (build plan s 14.7) ---------------------------------------------
from jaf.estimate import validate_spec, EST_PARAMS
from jaf.design import Hypothesis, Measurement
dv = illustrative_design(wI)
r36 = PRE.freeze(wI, dv)                                                  # values present when frozen
t36 = copy.deepcopy(r36); t36["values_present_at_freeze"] = []; t36["created"] = "2020-01-01T00:00:00"
c36 = PRE.check(t36, wI, dv)
check("36 review 2: editing a record's date or its list of values present at freezing is detected",
      bool(r36["values_present_at_freeze"]) and not c36["record_intact"] and not c36["consistent"])

r37 = PRE.freeze(wI, illustrative_design_skeleton(), base=Params())
c37a = PRE.check(r37, wI, illustrative_design_skeleton(), Params(alpha=0.5))
d37 = illustrative_design_skeleton(); d37.hypotheses.append(Hypothesis("EH1", "duplicate")); d37.measurement.append(Measurement("H1", 0.9))
wb37 = load_workbook(io.BytesIO(write_template(wI, design=d37))); bio = io.BytesIO(); wb37.save(bio)
w37, d37r, i37 = read_workbook(bio.getvalue()); m37 = " | ".join(i["message"] for i in i37)
check("37 review 3 and 4: a changed model setting is a deviation; duplicate hypothesis IDs and measurement rows are rejected",
      not c37a["consistent"] and "alpha (0.6 then, 0.5 now)" in " ".join(c37a["deviations"])
      and w37 is None and "used more than once" in m37 and "described more than once" in m37)

class Linear(Model):
    """Stand-in model with linear moments: targets see theta0 + beta + sigma only; the holdout sees theta0 - beta."""
    def moments(self, values, obs):
        t, b, g = values["theta0"], values["beta"], values["sigma"]
        return np.array([(t + b + g) * (i + 1) if o.role == "target" else t - b for i, o in enumerate(obs)], float)
    def converged(self, values):
        return True
    def approximate(self, values):
        return None
d38 = Design(observations=[Observation(f"T{i}", "exit_share", "", "baseline", sd=0.1, value=6.0 * (i + 1)) for i in range(3)]
             + [Observation("V1", "exit_share", "", "baseline", sd=0.1, value=0.0, role="holdout")],
             params=[ParamSpec("theta0", True, 3.0, 0.0, 10.0), ParamSpec("beta", True, 2.0, 0.0, 10.0), ParamSpec("sigma", True, 1.0, 0.0, 10.0)], starts=1)
m38 = Linear(wI, d38)
id38 = identification(m38)
e38 = estimate(m38, override=True); v38 = validate(m38, e38)
check("38 review 1 and 6: with rank deficiency 2 every parameter is 'not identified'; an undetermined holdout cannot be scored",
      [p_["status"] for p_ in id38["params"]] == ["not identified"] * 3 and id38["rank"] == 1
      and v38["unscorable"] == 1 and v38["n"] == 0 and "z" not in v38["rows"][0], f"rank {id38['rank']}")

import importlib.util, sys as _sys
spec_ = importlib.util.spec_from_file_location("srv", "app/server.py"); srv = importlib.util.module_from_spec(spec_); spec_.loader.exec_module(srv)
txt39 = json.dumps(srv._finite({"a": float("inf"), "b": [1.0, float("nan"), {"c": np.float64("-inf")}], "d": np.array([1.0, np.inf])}), allow_nan=False)
check("39 review 5: non-finite numbers are sent to the browser as null, never as Infinity or NaN",
      txt39 == '{"a": null, "b": [1.0, null, {"c": null}], "d": [1.0, null]}', txt39)

e40 = validate_spec([ParamSpec("alpha", False, 1.3), ParamSpec("sigma", False, -1.0), ParamSpec("theta0", True, 9.0, 5.0, 5.0),
                     ParamSpec("beta", True, 30.0, 40.0, 20.0), ParamSpec("beta", False, 25.0), ParamSpec("exit_weight", False, 0.6, judgement_low=0.3)])
check("40 review 7: parameter values outside what the model admits, reversed or empty bounds and duplicates are rejected",
      len(e40) >= 6 and any("alpha" in x and "outside" in x for x in e40) and any("sigma" in x for x in e40)
      and any("theta0" in x and "below the upper" in x for x in e40) and any("more than once" in x for x in e40)
      and any("both ends" in x for x in e40), f"{len(e40)} errors")

print(f"\n{sum(results)}/{len(results)} passed")
raise SystemExit(len(results) - sum(results))
