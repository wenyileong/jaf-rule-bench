# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0
# Copyright (c) 2026 Wen Yi LEONG
"""Excel coding template: write, read and validate (build plan s 12)."""
from __future__ import annotations
import io, datetime
from dataclasses import fields
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter
from .world import (World, Jurisdiction, Link, Settings, Treaty, LIABILITY, COLLECTION, KINDS, RATE_MODES, LINK_TYPES,
                    INTERMEDIARY_TESTS, illustrative_world, CANDIDATE_SAMPLE)

YES_NO = ("yes", "no")

# column key, header shown to the user, help text, allowed values (or None), required
J_COLS = [
    ("code", "Code", "Short unique code, e.g. UK. Used in the Personal links sheet.", None, True),
    ("name", "Name", "Full name of the jurisdiction.", None, True),
    ("kind", "Kind", "economy, hub, or rest_of_world.", KINDS, True),
    ("holder_base", "Holder base (relative)", "Relative size of the digital-asset holder base resident here. Only ratios matter.", None, True),
    ("exchange_presence", "Exchange presence (relative)", "How attractive this place is for holding assets with an exchange (1 = average). 0 if none.", None, True),
    ("liability_rule", "Liability rule", "Which connecting factor this jurisdiction taxes on. See the Lists sheet.", tuple(LIABILITY), True),
    ("intermediary_test", "Intermediary located by", "Only if the liability rule is intermediary: seat (incorporation or licence) or base (management and operations).", tuple(INTERMEDIARY_TESTS), False),
    ("collection", "Collection", "self_assessment, or domestic_withholding if domestic intermediaries withhold or collect the tax.", tuple(COLLECTION), True),
    ("offshore_incorporation_in", "Exchanges here incorporated offshore in (code)", "Typical place of incorporation of exchanges operating here, if elsewhere. Leave blank if they are incorporated here.", None, False),
    ("offshore_incorporation_share", "Share incorporated offshore (0 to 1)", "Share of exchanges operating here that are incorporated in that place.", None, False),
    ("domestic_exchanges_permitted", "Domestic exchanges permitted", "no if exchanges are prohibited from operating here.", YES_NO, True),
    ("visibility_domestic_custodial", "Visibility: domestic exchange holdings (0 to 1)", "Share of holdings at domestic exchanges the tax authority can see (domestic reporting).", None, True),
    ("visibility_self_hosted", "Visibility: self-hosted and unreported (0 to 1)", "Share of self-hosted holdings, or foreign holdings outside exchange of information, the authority can see.", None, True),
    ("info_exchange_from", "Automatic exchange from (year)", "First year of automatic exchange of crypto-asset information (CARF or DAC8). Leave blank if none.", None, False),
    ("residence_topup", "Residence top-up on untaxed holdings (0 to 1)", "Share of its own rate applied to residents' holdings that no rule assigns. 0 if none.", None, True),
    ("rate_mode", "Rate mode", "fixed: use the rate in the next column. competitive: the model sets the rate by competition.", RATE_MODES, True),
    ("effective_rate", "Effective rate on returns (0 to 1)", "Required if rate mode is fixed. The effective tax rate on the return from holdings, not the headline statutory rate; see Read me.", None, False),
    ("in_force_from", "In force from", "Date the coded rules took effect (YYYY-MM-DD).", None, False),
    ("in_force_until", "In force until", "Date the coded rules ceased, if any.", None, False),
    ("source", "Source (OSCOLA)", "Legislation, guidance or case relied on for this row.", None, False),
    ("verified", "Verified", "yes once the row has been checked against the source.", YES_NO, True),
    ("notes", "Notes", "Anything a reader should know about this row.", None, False),
]
L_COLS = [
    ("from_code", "Holders resident in (code)", "", None, True),
    ("to_code", "Also connected to (code)", "", None, True),
    ("link_type", "Link type", "dual_residence or citizenship.", LINK_TYPES, True),
    ("share", "Share of residents (0 to 1)", "", None, True),
    ("source", "Source (OSCOLA)", "", None, False),
    ("verified", "Verified", "", YES_NO, True),
]
T_COLS = [
    ("relieving", "Relieving state (code)", "The state that gives relief, usually the residence state.", None, True),
    ("taxing", "Taxing state (code)", "The state whose prior tax is relieved.", None, True),
    ("relief", "Relief (0 to 1)", "Share of the double tax relieved: 1 for a full credit or exemption, 0 for none.", None, True),
    ("both_directions", "Both directions", "yes if the same relief applies with the roles reversed.", YES_NO, True),
    ("source", "Source (OSCOLA)", "Treaty article or domestic relief provision.", None, False),
    ("verified", "Verified", "", YES_NO, True),
]
S_ROWS = [
    ("label", "Label for this world", "A name for the coded world."),
    ("scenario_year", "Scenario year", "Exchange of information is active for jurisdictions whose start year is on or before this year."),
    ("visibility_exchanged", "Visibility of exchanged information (0 to 1)", "Quality of information received through automatic exchange."),
    ("mobility_custodial", "Ease of holding with a foreign exchange", "Scaled by the destination's exchange presence. Default 0.40."),
    ("mobility_self_hosted", "Ease of placing self-hosted keys abroad", "Default 0.90."),
    ("treaty_relief", "Treaty relief for double taxation (0 to 1)", "Share of double tax relieved where two states claim. Default 0.75."),
    ("illustrative", "Illustrative world", "yes if the rows are hypothetical. Results from such a world must not be presented as describing real states."),
]

README = [
    "JAF Rule Bench: coding template",
    "",
    "What this is",
    "This workbook describes a set of jurisdictions for the model: which connecting factor each taxes on (liability),",
    "who collects the tax (collection), and what its tax authority can see (observation). Fill one row per jurisdiction",
    "in the Jurisdictions sheet, add any dual residence or citizenship links in the Personal links sheet, and set the",
    "global values in the Settings sheet. Then load the file on the 'Coded world' tab.",
    "",
    "How to fill it",
    "1. Use the drop-down lists where they appear. The Lists sheet explains every allowed value.",
    "2. Cite a source for every row in OSCOLA form and set Verified to yes only once checked.",
    "   The tool will run with unverified rows but labels every result that depends on them.",
    "3. Record the dates the coded rules were in force. Code the law as at the Scenario year.",
    "4. Effective rate: the model works with the effective tax rate on the return from holdings, not the headline",
    "   statutory rate. For a flat tax on all realised gains, with no loss relief, it is close to the statutory rate.",
    "   It is lower where gains are exempt after a holding period, taxed only on realisation (deferral), or reduced",
    "   by allowances; it can be higher where losses cannot be set off (as under India's 2022 regime) or where a",
    "   transaction tax applies to turnover rather than gains. Record the conversion and its assumptions in Notes.",
    "   The user guide (docs/USER_GUIDE.md, 'From statutory to effective rates') sets out a worked method.",
    "5. Include low-tax hubs and a 'Rest of world' row, so that holdings leaving the sample have somewhere to go.",
    "",
    "Modelling assumptions to keep in mind",
    "- Co-location: a holding is housed in one place, and its location-based connecting factors follow it, except",
    "  that an exchange may operate in one state (its base) and be incorporated in another (its seat); see below.",
    "- Where several states claim the same holding, relief runs in this order: the state where the exchange",
    "  operates (or where the holding is housed) taxes first, then the state of incorporation, then a co-signing",
    "  state, then the residence state, then a citizenship state.",
    "- Withholding by domestic intermediaries is treated as full visibility of domestic exchange holdings. The model",
    "  assumes the tax falls on the holder whoever collects it; who bears a withholding tax in practice is not modelled.",
    "- Exchanges have an operational base (where they are run) and may have a legal seat elsewhere (where they are",
    "  incorporated or licensed). Code, for each state, where exchanges operating there are usually incorporated.",
    "  Registration follows the seat; infrastructure, key control and validation follow the base; a state that taxes",
    "  by intermediary location applies its own test (seat or base).",
    "- Treaty relief: the Treaties sheet records relief by pair and direction. Pairs not listed use the default in",
    "  Settings.",
    "",
    "The rows in an 'Illustrative' workbook describe hypothetical states and must not be cited as law.",
    "",
    "Licence",
    "This template is part of JAF Rule Bench, Copyright (c) 2026 Wen Yi LEONG, and is licensed under the Creative",
    "Commons Attribution-NonCommercial 4.0 International licence (CC BY-NC 4.0):",
    "https://creativecommons.org/licenses/by-nc/4.0/. The entries you add to it are your own.",
    "",
    "Data for testing the model (optional; build plan s 14)",
    "The sheets Episodes, Data, Measurement, Estimation and Hypotheses hold an empirical design. They are needed only",
    "to estimate the model's behavioural parameters from observed facts and to test its predictions (the 'Test against",
    "data' tab). Leave them empty to use the tool as before.",
    "- Episodes: versions of the coded law to compare. 'baseline' is the coded world as loaded. Each other episode is a",
    "  list of changes, one per row: the jurisdiction code (or 'settings'), the field to change and its new value.",
    "  A single reform (for example India's 2022 regime) is one episode that changes one state's rate and collection.",
    "  Staggered adoption of automatic exchange (for example the CARF waves) is one episode per wave, each setting the",
    "  scenario year (and, if needed, a state's first year of exchange).",
    "- Data: one row per observed fact. Choose the quantity from the drop-down list (the Quantities sheet explains what",
    "  each measures and how data record it), the jurisdiction, the episode and, for a change, the episode it is",
    "  compared with. Give the standard error of the value: it says how far the value could be off, and the fit is",
    "  weighted by it. Mark each row target (used to estimate) or holdout (kept back to test the predictions).",
    "  Leave values blank, fill everything else, and freeze the design on the tab before looking at the data: that",
    "  is the pre-registration.",
    "- Measurement: for exchange data, the share of reported volume that is wash trading at exchanges based in each",
    "  place, and genuine turnover relative to elsewhere (1 if unknown). Volume-based quantities are adjusted for both.",
    "- Estimation: which behavioural parameters to estimate (free = yes, with a start value and bounds) and which to fix",
    "  by judgement (free = no, with a value and, optionally, a plausible range used to widen prediction bands).",
    "- Hypotheses: the claims the data are to test, and what result would support or refute each.",
]
README_HEADS = ("What this is", "How to fill it", "Modelling assumptions to keep in mind", "Licence",
                "Data for testing the model (optional; build plan s 14)")


def _cell(v):
    if isinstance(v, bool):
        return "yes" if v else "no"
    return v


def write_template(world: World | None = None, kind: str = "blank", design=None) -> bytes:
    """kind: 'illustrative' (hypothetical states, filled) or 'blank' (candidate sample, legal fields empty).
    design: the empirical design to write in the data sheets (default: the illustrative design for an
    illustrative workbook, an empty design otherwise)."""
    if world is None and kind == "illustrative":
        world = illustrative_world()
        if design is None:
            from .examples import illustrative_design
            design = illustrative_design(world)
    wb = Workbook()
    head = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="3A2FD6")
    ws = wb.active; ws.title = "Read me"
    for r, line in enumerate(README, 1):
        ws.cell(r, 1, line)
        if line in ("JAF Rule Bench: coding template",):
            ws.cell(r, 1).font = Font(bold=True, size=14)
        elif line in README_HEADS:
            ws.cell(r, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 120

    def sheet(title, cols, rows):
        sh = wb.create_sheet(title)
        for ci, (key, label, helptext, allowed, req) in enumerate(cols, 1):
            c = sh.cell(1, ci, label + (" *" if req else "")); c.font = head; c.fill = fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
            sh.cell(2, ci, helptext).alignment = Alignment(wrap_text=True, vertical="top")
            sh.cell(2, ci).font = Font(italic=True, color="666666", size=9)
            sh.column_dimensions[get_column_letter(ci)].width = max(14, min(40, len(label) + 4))
            if allowed:
                dv = DataValidation(type="list", formula1='"' + ",".join(allowed) + '"', allow_blank=not req)
                sh.add_data_validation(dv); dv.add(f"{get_column_letter(ci)}3:{get_column_letter(ci)}200")
        sh.row_dimensions[1].height = 45; sh.row_dimensions[2].height = 60
        sh.freeze_panes = "C3"
        for ri, row in enumerate(rows, 3):
            for ci, (key, *_rest) in enumerate(cols, 1):
                v = row.get(key)
                if v is not None and v != "":
                    sh.cell(ri, ci, _cell(v))
        return sh

    if world is not None:
        def blank(j, name):
            return ((name == "intermediary_test" and j.liability_rule != "intermediary")
                    or (name == "offshore_incorporation_share" and not j.offshore_incorporation_in))
        jrows = [{f.name: ("" if blank(j, f.name) else getattr(j, f.name)) for f in fields(Jurisdiction)} for j in world.jurisdictions]
        lrows = [{f.name: getattr(l, f.name) for f in fields(Link)} for l in world.links]
        trows = [{f.name: getattr(t, f.name) for f in fields(Treaty)} for t in world.treaties]
        settings = world.settings
    else:
        jrows = [dict(code=c, name=n, kind=k, verified=False) for c, n, k in CANDIDATE_SAMPLE]
        lrows, trows = [], []
        settings = Settings(label="Real-world sample (to be coded)", illustrative=False)
    sheet("Jurisdictions", J_COLS, jrows)
    sheet("Personal links", L_COLS, lrows)
    sheet("Treaties", T_COLS, trows)
    st = wb.create_sheet("Settings")
    for ci, h in enumerate(("Setting", "Value", "Explanation"), 1):
        c = st.cell(1, ci, h); c.font = head; c.fill = fill
    for ri, (key, label, helptext) in enumerate(S_ROWS, 2):
        st.cell(ri, 1, label); st.cell(ri, 2, _cell(getattr(settings, key))); st.cell(ri, 3, helptext)
        st.cell(ri, 4, key).font = Font(color="BBBBBB", size=8)
    st.column_dimensions["A"].width = 44; st.column_dimensions["B"].width = 34; st.column_dimensions["C"].width = 90
    _write_design(wb, sheet, st, design, head, fill)
    ls = wb.create_sheet("Lists")
    r = 1
    for title, d in (("Liability rule", LIABILITY), ("Collection", COLLECTION)):
        ls.cell(r, 1, title).font = Font(bold=True); r += 1
        for k, v in d.items():
            ls.cell(r, 1, k); ls.cell(r, 2, v); r += 1
        r += 1
    for title, vals in (("Kind", KINDS), ("Rate mode", RATE_MODES), ("Link type", LINK_TYPES), ("Intermediary located by", tuple(INTERMEDIARY_TESTS))):
        ls.cell(r, 1, title).font = Font(bold=True); r += 1
        for v in vals:
            ls.cell(r, 1, v); r += 1
        r += 1
    for title, vals in (("Episode fields (a jurisdiction)", EPISODE_J_FIELDS), ("Episode fields (settings)", EPISODE_S_FIELDS),
                        ("Comparison", COMPARISONS), ("Role", ROLES)):
        ls.cell(r, 1, title).font = Font(bold=True); r += 1
        for v in vals:
            ls.cell(r, 1, v); r += 1
        r += 1
    ls.cell(r, 1, "Parameters that can be estimated").font = Font(bold=True); r += 1
    for k, v in EST_PARAMS.items():
        ls.cell(r, 1, k); ls.cell(r, 2, f"{v['plain']} ({v['tech']}); default range {v['lo']} to {v['hi']}"); r += 1
    ls.column_dimensions["A"].width = 30; ls.column_dimensions["B"].width = 90
    bio = io.BytesIO(); wb.save(bio); return bio.getvalue()


# ---- reading and validation ----------------------------------------------------------
def _norm(v):
    if v is None:
        return ""
    if isinstance(v, str):
        return v.strip()
    return v


def _yes(v):
    return str(v).strip().lower() in ("yes", "y", "true", "1")


def read_template(data: bytes):
    """Returns (world or None, issues). issues: list of dict(level, where, message)."""
    issues = []
    def err(where, msg): issues.append(dict(level="error", where=where, message=msg))
    def warn(where, msg): issues.append(dict(level="warning", where=where, message=msg))
    try:
        wb = load_workbook(io.BytesIO(data), data_only=True)
    except Exception:
        err("File", "This file could not be opened as an Excel workbook (.xlsx).")
        return None, issues
    for name in ("Jurisdictions", "Personal links", "Settings"):
        if name not in wb.sheetnames:
            err("File", f"The sheet '{name}' is missing. Start from a template downloaded from the tool.")
    if any(i["level"] == "error" for i in issues):
        return None, issues

    absent = set()

    def rows(sheet, cols):
        sh = wb[sheet]
        hdr = [str(_norm(c.value)).rstrip(" *") for c in sh[1]]
        idx = {}
        for key, label, _h, _a, req in cols:
            if label in hdr:
                idx[key] = hdr.index(label)
            elif req:
                err(sheet, f"Column '{label}' is missing.")
            else:
                absent.add((sheet, key))
                warn(sheet, f"Column '{label}' is missing (template from an earlier version); defaults are used.")
        out = []
        for r in sh.iter_rows(min_row=3, values_only=True):
            if r is None or all(_norm(x) in ("", None) for x in r):
                continue
            out.append((dict((k, _norm(r[i]) if i < len(r) else "") for k, i in idx.items())))
        return out

    jr = rows("Jurisdictions", J_COLS); lr = rows("Personal links", L_COLS)
    tr = rows("Treaties", T_COLS) if "Treaties" in wb.sheetnames else []
    if any(i["level"] == "error" for i in issues):
        return None, issues

    J, seen = [], set()
    for n, r in enumerate(jr, 3):
        where = f"Jurisdictions, row {n}" + (f" ({r.get('code')})" if r.get("code") else "")
        code = str(r.get("code", "")).strip()
        if not code:
            err(where, "Code is empty."); continue
        if code in seen:
            err(where, f"Code '{code}' is used more than once.")
        seen.add(code)
        def num(key, lo=0.0, hi=None, required=True):
            v = r.get(key, "")
            if v == "" or v is None:
                if required:
                    err(where, f"'{dict((k, l) for k, l, *_ in J_COLS)[key]}' is empty.")
                return None
            try:
                x = float(v)
            except (TypeError, ValueError):
                err(where, f"'{dict((k, l) for k, l, *_ in J_COLS)[key]}' must be a number (found '{v}')."); return None
            if x < lo or (hi is not None and x > hi):
                err(where, f"'{dict((k, l) for k, l, *_ in J_COLS)[key]}' must be between {lo} and {hi if hi is not None else 'any'} (found {x}).")
            return x
        def choice(key, allowed):
            v = str(r.get(key, "")).strip()
            if v not in allowed:
                err(where, f"'{dict((k, l) for k, l, *_ in J_COLS)[key]}' must be one of: {', '.join(allowed)} (found '{v or 'blank'}').")
                return allowed[0]
            return v
        kind = choice("kind", KINDS); rule = choice("liability_rule", tuple(LIABILITY)); coll = choice("collection", tuple(COLLECTION))
        itest = str(r.get("intermediary_test", "") or "").strip()
        if itest and itest not in INTERMEDIARY_TESTS:
            err(where, f"'Intermediary located by' must be seat or base (found '{itest}').")
        if rule == "intermediary" and not itest:
            if ("Jurisdictions", "intermediary_test") in absent:
                warn(where, "Assumed to locate intermediaries by legal seat, because the template has no 'Intermediary located by' column. Add the column to state the test.")
            else:
                err(where, "This state taxes by intermediary location, so 'Intermediary located by' must be seat or base.")
        off = str(r.get("offshore_incorporation_in", "") or "").strip()
        offsh = r.get("offshore_incorporation_share", "")
        offshare = 0.0
        if off:
            try:
                offshare = float(offsh)
                if not 0 <= offshare <= 1:
                    err(where, "'Share incorporated offshore' must be between 0 and 1.")
            except (TypeError, ValueError):
                err(where, "'Share incorporated offshore' is needed when an offshore place of incorporation is given.")
        elif offsh not in ("", None) and str(offsh).strip() not in ("0", "0.0"):
            warn(where, "A share incorporated offshore is given but no place of incorporation; it is ignored.")
        mode = choice("rate_mode", RATE_MODES)
        permitted = str(r.get("domestic_exchanges_permitted", "")).strip().lower()
        if permitted not in YES_NO:
            err(where, "'Domestic exchanges permitted' must be yes or no.")
        hb = num("holder_base", 0.0); ep = num("exchange_presence", 0.0)
        vd = num("visibility_domestic_custodial", 0.0, 1.0); vs = num("visibility_self_hosted", 0.0, 1.0)
        top = num("residence_topup", 0.0, 1.0)
        rate = num("effective_rate", 0.0, 1.0, required=(mode == "fixed"))
        if mode == "fixed" and rate is None:
            pass
        yr = r.get("info_exchange_from", "")
        year = None
        if yr not in ("", None):
            try:
                year = int(float(yr))
                if not 1990 <= year <= 2100:
                    err(where, f"'Automatic exchange from' must be a year (found {yr}).")
            except (TypeError, ValueError):
                err(where, f"'Automatic exchange from' must be a year or blank (found '{yr}').")
        if hb is not None and hb == 0 and kind != "hub":
            warn(where, "Holder base is 0: no holders are resident here.")
        if permitted == "no" and ep and ep > 0:
            warn(where, "Domestic exchanges are prohibited, so exchange presence is ignored.")
        verified = _yes(r.get("verified", ""))
        src = str(r.get("source", "")).strip()
        if not src:
            warn(where, "No source cited.")
        if not verified:
            warn(where, "Not yet verified. Results that depend on this row will be labelled.")
        J.append(Jurisdiction(code=code, name=str(r.get("name") or code), kind=kind, holder_base=hb or 0.0,
                              exchange_presence=ep or 0.0, liability_rule=rule, collection=coll,
                              domestic_exchanges_permitted=(permitted != "no"),
                              visibility_domestic_custodial=vd if vd is not None else 0.0,
                              visibility_self_hosted=vs if vs is not None else 0.0, info_exchange_from=year,
                              residence_topup=top or 0.0, rate_mode=mode, effective_rate=rate,
                              in_force_from=str(r.get("in_force_from") or ""), in_force_until=str(r.get("in_force_until") or ""),
                              source=src, verified=verified, notes=str(r.get("notes") or ""),
                              intermediary_test=itest or "seat", offshore_incorporation_in=off, offshore_incorporation_share=offshare))
    if len(J) < 2:
        err("Jurisdictions", "At least two jurisdictions are needed.")
    if J and not any(j.kind == "rest_of_world" for j in J):
        warn("Jurisdictions", "There is no 'rest_of_world' row, so holdings leaving the sample have nowhere else to go.")
    if J and all(j.liability_rule == "none" for j in J):
        warn("Jurisdictions", "No jurisdiction assigns anything; every holding is in the gap.")
    if J and sum(j.holder_base for j in J) <= 0:
        err("Jurisdictions", "All holder bases are 0.")

    L = []
    codes = {j.code for j in J}
    for n, r in enumerate(lr, 3):
        where = f"Personal links, row {n}"
        a, b = str(r.get("from_code", "")).strip(), str(r.get("to_code", "")).strip()
        if a not in codes or b not in codes:
            err(where, f"Codes must match the Jurisdictions sheet (found '{a}' and '{b}')."); continue
        if a == b:
            err(where, "A link must connect two different jurisdictions."); continue
        t = str(r.get("link_type", "")).strip()
        if t not in LINK_TYPES:
            err(where, f"Link type must be dual_residence or citizenship (found '{t or 'blank'}')."); continue
        try:
            s = float(r.get("share", ""))
        except (TypeError, ValueError):
            err(where, "Share must be a number between 0 and 1."); continue
        if not 0 <= s <= 1:
            err(where, "Share must be between 0 and 1."); continue
        v = _yes(r.get("verified", ""))
        if not v:
            warn(where, "Not yet verified.")
        L.append(Link(a, b, t, s, str(r.get("source") or ""), v))
    for j in J:
        if j.offshore_incorporation_in and j.offshore_incorporation_in not in codes:
            err(f"Jurisdictions ({j.code})", f"Offshore place of incorporation '{j.offshore_incorporation_in}' is not a code in the Jurisdictions sheet.")
        elif j.offshore_incorporation_in == j.code:
            err(f"Jurisdictions ({j.code})", "The offshore place of incorporation must be a different jurisdiction.")
    T, seen_pairs = [], set()
    for n, r in enumerate(tr, 3):
        where = f"Treaties, row {n}"
        q, p_ = str(r.get("relieving", "")).strip(), str(r.get("taxing", "")).strip()
        if q not in codes or p_ not in codes or q == p_:
            err(where, f"Codes must be two different jurisdictions from the Jurisdictions sheet (found '{q}' and '{p_}')."); continue
        try:
            rv = float(r.get("relief", ""))
        except (TypeError, ValueError):
            err(where, "Relief must be a number between 0 and 1."); continue
        if not 0 <= rv <= 1:
            err(where, f"Relief must be between 0 and 1 (found {rv})."); continue
        both = _yes(r.get("both_directions", ""))
        for pair in ([(q, p_), (p_, q)] if both else [(q, p_)]):
            if pair in seen_pairs:
                warn(where, f"Relief for {pair[0]} relieving {pair[1]} is set more than once; the last row applies.")
            seen_pairs.add(pair)
        v = _yes(r.get("verified", ""))
        if not v:
            warn(where, "Not yet verified.")
        T.append(Treaty(q, p_, rv, both, str(r.get("source") or ""), v))
    for code in codes:
        tot = sum(l.share for l in L if l.from_code == code)
        if tot > 1:
            err("Personal links", f"Links from {code} add up to more than 1 ({tot:.2f}).")

    st = wb["Settings"]; svals = {}
    for row in st.iter_rows(min_row=2, values_only=True):
        if row and len(row) >= 4 and row[3]:
            svals[str(row[3]).strip()] = row[1]
    s = Settings()
    for key, label, _ in S_ROWS:
        if key not in svals or svals[key] in (None, ""):
            continue
        v = svals[key]
        try:
            if key == "label":
                s.label = str(v)
            elif key == "illustrative":
                s.illustrative = _yes(v)
            elif key == "scenario_year":
                s.scenario_year = int(float(v))
            else:
                x = float(v)
                if key in ("visibility_exchanged", "treaty_relief") and not 0 <= x <= 1:
                    err("Settings", f"'{label}' must be between 0 and 1."); continue
                setattr(s, key, x)
        except (TypeError, ValueError):
            err("Settings", f"'{label}' has an invalid value ('{v}').")
    if any(i["level"] == "error" for i in issues):
        return None, issues
    return World(J, L, s, T), issues


# ---- empirical design sheets (build plan s 14, B1) ---------------------------------------
from .design import (Design, Episode, Change, Observation, Measurement, ParamSpec, Hypothesis, BASELINE, ROLES,
                     COMPARISONS, EPISODE_J_FIELDS, EPISODE_S_FIELDS, coerce)
from .measure import QUANTITIES
from .estimate import EST_PARAMS, validate_spec

E_COLS = [
    ("episode", "Episode", "Name of the episode, e.g. reform_2022 or wave_2027. Repeat it on one row per change. Not 'baseline'.", None, True),
    ("target", "Applies to (code or settings)", "Jurisdiction code, or settings for scenario year, visibility of exchanged information or treaty relief.", None, True),
    ("field", "Field to change", "See the Lists sheet for the fields an episode may change.", tuple(EPISODE_J_FIELDS + EPISODE_S_FIELDS), True),
    ("value", "New value", "Value in the episode. Use yes or no, a number, a rule from the Lists sheet, or a year.", None, False),
    ("description", "Description of the episode", "Plain description (needed once per episode).", None, False),
    ("source", "Source (OSCOLA)", "Law or measure that makes the change.", None, False),
    ("verified", "Verified", "", YES_NO, True),
]
D_COLS = [
    ("id", "ID", "Short unique identifier, e.g. IN1.", None, True),
    ("quantity", "Quantity", "What is observed. See the Quantities sheet.", tuple(QUANTITIES), True),
    ("code", "Jurisdiction (code)", "The jurisdiction the quantity refers to. Blank only for exit_share.", None, False),
    ("episode", "Episode", "baseline or an episode from the Episodes sheet.", None, True),
    ("compared_with", "Compared with episode", "For a change: the episode it is measured against. Blank for a level.", None, False),
    ("comparison", "Comparison", "ratio (episode / compared) or difference (episode minus compared).", COMPARISONS, False),
    ("value", "Value", "Observed value. Leave blank until the design has been frozen (pre-registration).", None, False),
    ("sd", "Standard error", "How far the value could be off, in the same units. Required.", None, True),
    ("role", "Role", "target (used to estimate) or holdout (used only to test predictions).", ROLES, True),
    ("unit", "Unit", "What the value measures, e.g. share of volume, ratio.", None, False),
    ("source", "Source", "Data source (OSCOLA or dataset citation).", None, False),
    ("accessed", "Date accessed", "", None, False),
    ("quality", "Quality (A, B or C)", "A: official or audited; B: reputable industry or academic estimate; C: indicative.", ("A", "B", "C"), False),
    ("verified", "Verified", "", YES_NO, True),
    ("notes", "Notes", "Construction of the value, adjustments made.", None, False),
]
M_COLS = [
    ("code", "Exchanges based in (code)", "", None, True),
    ("wash_share", "Wash-trading share of reported volume (0 to 1)", "0 if none or unknown.", None, False),
    ("turnover", "Genuine turnover relative to elsewhere", "1 if unknown.", None, False),
    ("source", "Source", "", None, False),
    ("verified", "Verified", "", YES_NO, True),
]
P_COLS = [
    ("name", "Parameter", "See the Lists sheet.", tuple(EST_PARAMS), True),
    ("free", "Estimate it?", "yes: estimated from the targets. no: fixed at the value given (or the interface value if blank).", YES_NO, True),
    ("value", "Start or fixed value", "", None, False),
    ("lower", "Lower bound", "For estimated parameters (blank: the tool's default range).", None, False),
    ("upper", "Upper bound", "", None, False),
    ("judgement_low", "Judgement range: low", "For fixed parameters: low end of the plausible range (widens prediction bands).", None, False),
    ("judgement_high", "Judgement range: high", "", None, False),
    ("notes", "Notes", "Why it is estimated or fixed, and the basis for any fixed value.", None, False),
]
H_COLS = [
    ("id", "ID", "", None, True),
    ("statement", "Hypothesis", "", None, True),
    ("test", "Test", "Which observations or results would support or refute it.", None, False),
    ("source", "Source", "Where the hypothesis comes from (for example a proposition of the paper).", None, False),
]
DESIGN_ROWS = [
    ("design_label", "Label for the empirical design", "A name for the design."),
    ("estimation_concept", "Equilibrium concept for estimation", "a (states treat holders' reactions as given) or b (states also foresee escape routes)."),
    ("estimation_starts", "Number of starting points", "The search is repeated from this many points; more guards against a local minimum."),
    ("estimation_seed", "Random seed for starting points", ""),
    ("synthetic", "Synthetic data", "yes if the values were generated by the tool rather than observed."),
]
DESIGN_SHEETS = ("Episodes", "Data", "Measurement", "Estimation", "Hypotheses")


def _write_design(wb, sheet, st, design, head, fill):
    d = design if design is not None else Design(params=[ParamSpec(n, False) for n in EST_PARAMS])
    erows = []
    for e in d.episodes:
        for n, c in enumerate(e.changes):
            erows.append(dict(episode=e.name, target=c.target, field=c.field, value=c.value,
                              description=e.description if n == 0 else "", source=c.source, verified=c.verified))
        if not e.changes:
            erows.append(dict(episode=e.name, description=e.description))
    sheet("Episodes", E_COLS, erows)
    sheet("Data", D_COLS, [{f: getattr(o, f) for f in (c[0] for c in D_COLS)} for o in d.observations])
    sheet("Measurement", M_COLS, [{f: getattr(m, f) for f in (c[0] for c in M_COLS)} for m in d.measurement])
    sheet("Estimation", P_COLS, [{f: getattr(p, f) for f in (c[0] for c in P_COLS)} for p in d.params])
    sheet("Hypotheses", H_COLS, [{f: getattr(h, f) for f in (c[0] for c in H_COLS)} for h in d.hypotheses])
    vals = dict(design_label=d.label, estimation_concept=d.concept, estimation_starts=d.starts, estimation_seed=d.seed, synthetic=d.synthetic)
    r0 = st.max_row + 2
    st.cell(r0 - 1, 1, "Empirical design").font = Font(bold=True)
    for ri, (key, label, helptext) in enumerate(DESIGN_ROWS, r0):
        st.cell(ri, 1, label); st.cell(ri, 2, _cell(vals[key])); st.cell(ri, 3, helptext)
        st.cell(ri, 4, key).font = Font(color="BBBBBB", size=8)
    q = wb.create_sheet("Quantities")
    for ci, h in enumerate(("Quantity", "What it measures", "Technical definition", "Recorded as", "Level usable", "Use"), 1):
        c = q.cell(1, ci, h); c.font = head; c.fill = fill
    for ri, (k, v) in enumerate(QUANTITIES.items(), 2):
        for ci, x in enumerate((k, v["plain"], v["tech"], "reported volume" if v["basis"] == "volume" else "holdings",
                                "yes" if v["level"] else "no: only as a ratio or difference between episodes", v["use"]), 1):
            q.cell(ri, ci, x).alignment = Alignment(wrap_text=True, vertical="top")
    for col, wdt in zip("ABCDEF", (34, 50, 60, 16, 22, 70)):
        q.column_dimensions[col].width = wdt


def _rows_generic(wb, sheet, cols, issues, absent=None):
    sh = wb[sheet]
    hdr = [str(_norm(c.value)).rstrip(" *") for c in sh[1]]
    idx = {}
    for key, label, _h, _a, req in cols:
        if label in hdr:
            idx[key] = hdr.index(label)
        elif req:
            issues.append(dict(level="error", where=sheet, message=f"Column '{label}' is missing."))
    out = []
    for n, r in enumerate(sh.iter_rows(min_row=3, values_only=True), 3):
        if r is None or all(_norm(x) in ("", None) for x in r):
            continue
        out.append((n, dict((k, _norm(r[i]) if i < len(r) else "") for k, i in idx.items())))
    return out


def _num(v):
    if v in ("", None):
        return None
    return float(v)


def read_design(wb, world: World):
    """Read the empirical design sheets. Returns (design or None, issues). None if the workbook has no design."""
    issues = []
    err = lambda w, m: issues.append(dict(level="error", where=w, message=m))
    warn = lambda w, m: issues.append(dict(level="warning", where=w, message=m))
    if not any(n in wb.sheetnames for n in DESIGN_SHEETS):
        return None, issues
    codes = set(world.codes) if world else set()
    rows = {n: (_rows_generic(wb, n, cols, issues) if n in wb.sheetnames else [])
            for n, cols in (("Episodes", E_COLS), ("Data", D_COLS), ("Measurement", M_COLS), ("Estimation", P_COLS), ("Hypotheses", H_COLS))}
    if any(i["level"] == "error" for i in issues):
        return None, issues
    if not any(rows[n] for n in ("Episodes", "Data", "Hypotheses")) and not any(_yes(r.get("free")) for _, r in rows["Estimation"]):
        return None, issues                                   # empty design sheets: nothing to read
    # episodes
    eps = {}
    for n, r in rows["Episodes"]:
        where = f"Episodes, row {n}"
        name = str(r.get("episode", "")).strip()
        if not name:
            err(where, "Episode name is empty."); continue
        if name == BASELINE:
            err(where, "'baseline' is the coded world itself and cannot be changed by an episode."); continue
        e = eps.setdefault(name, Episode(name))
        if r.get("description"):
            e.description = str(r["description"])
        tgt, fld = str(r.get("target", "")).strip(), str(r.get("field", "")).strip()
        if not tgt and not fld:
            continue
        if tgt != "settings" and tgt not in codes:
            err(where, f"'{tgt or 'blank'}' is neither a jurisdiction code nor 'settings'."); continue
        allowed = EPISODE_S_FIELDS if tgt == "settings" else EPISODE_J_FIELDS
        if fld not in allowed:
            err(where, f"An episode cannot change '{fld or 'blank'}' for {tgt}. Allowed: {', '.join(allowed)}."); continue
        try:
            val = coerce(fld, r.get("value"), for_settings=(tgt == "settings"))
        except (TypeError, ValueError) as ex:
            err(where, f"New value for {fld}: {ex if str(ex) else 'invalid'} (found '{r.get('value')}')."); continue
        if fld == "liability_rule" and val not in LIABILITY:
            err(where, f"Unknown liability rule '{val}'."); continue
        if fld == "collection" and val not in COLLECTION:
            err(where, f"Unknown collection '{val}'."); continue
        if fld == "rate_mode" and val not in RATE_MODES:
            err(where, f"Rate mode must be fixed or competitive (found '{val}')."); continue
        if fld == "intermediary_test" and val not in INTERMEDIARY_TESTS:
            err(where, "Intermediary test must be seat or base."); continue
        if fld == "offshore_incorporation_in" and val and val not in codes:
            err(where, f"'{val}' is not a jurisdiction code."); continue
        v = _yes(r.get("verified", ""))
        if not v:
            warn(where, "Not yet verified.")
        e.changes.append(Change(tgt, fld, val, str(r.get("source") or ""), v))
    for e in eps.values():
        if not e.changes:
            warn(f"Episodes ({e.name})", "This episode changes nothing: it is identical to the baseline.")
        if world is not None:
            from .design import apply_episode
            w2 = apply_episode(world, e)
            for j in w2.jurisdictions:
                if j.rate_mode == "fixed" and j.effective_rate is None:
                    err(f"Episodes ({e.name})", f"{j.code} has a fixed rate mode but no effective rate in this episode.")
    names = {BASELINE} | set(eps)
    # observations
    obs, seen = [], set()
    for n, r in rows["Data"]:
        where = f"Data, row {n}" + (f" ({r.get('id')})" if r.get("id") else "")
        oid = str(r.get("id", "")).strip()
        if not oid:
            err(where, "ID is empty."); continue
        if oid in seen:
            err(where, f"ID '{oid}' is used more than once.")
        seen.add(oid)
        q = str(r.get("quantity", "")).strip()
        if q not in QUANTITIES:
            err(where, f"Unknown quantity '{q or 'blank'}'. See the Quantities sheet."); continue
        code = str(r.get("code", "") or "").strip()
        if QUANTITIES[q]["code"] and code not in codes:
            err(where, f"This quantity needs a jurisdiction code from the Jurisdictions sheet (found '{code or 'blank'}')."); continue
        if not QUANTITIES[q]["code"] and code:
            warn(where, "This quantity refers to the whole world; the jurisdiction code is ignored."); code = ""
        ep = str(r.get("episode", "") or "").strip() or BASELINE
        if ep not in names:
            err(where, f"Episode '{ep}' is not defined in the Episodes sheet."); continue
        cw = str(r.get("compared_with", "") or "").strip()
        if cw and cw not in names:
            err(where, f"Episode '{cw}' (compared with) is not defined."); continue
        if cw == ep:
            err(where, "An observation cannot be compared with its own episode."); continue
        comp = str(r.get("comparison", "") or "").strip() or "ratio"
        if comp not in COMPARISONS:
            err(where, "Comparison must be ratio or difference."); continue
        if not QUANTITIES[q]["level"] and not cw:
            err(where, f"'{q}' has no observed counterpart in levels; give an episode to compare with."); continue
        try:
            val = _num(r.get("value"))
            sd = _num(r.get("sd"))
        except (TypeError, ValueError):
            err(where, "Value and standard error must be numbers."); continue
        if sd is None or sd <= 0:
            err(where, "A positive standard error is required."); continue
        role = str(r.get("role", "") or "").strip()
        if role not in ROLES:
            err(where, "Role must be target or holdout."); continue
        if val is not None and QUANTITIES[q]["level"] and not cw and QUANTITIES[q]["unit"].startswith("share") and not 0 <= val <= 1:
            warn(where, f"A share is expected to lie between 0 and 1 (found {val}).")
        v = _yes(r.get("verified", ""))
        if val is not None and not v:
            warn(where, "Not yet verified.")
        if val is not None and not str(r.get("source") or "").strip():
            warn(where, "No source given for this value.")
        obs.append(Observation(oid, q, code, ep, cw, comp, val, sd, role, str(r.get("unit") or ""), str(r.get("source") or ""),
                               str(r.get("accessed") or ""), str(r.get("quality") or ""), v, str(r.get("notes") or "")))
    # measurement
    meas = []
    for n, r in rows["Measurement"]:
        where = f"Measurement, row {n}"
        code = str(r.get("code", "")).strip()
        if code not in codes:
            err(where, f"'{code or 'blank'}' is not a jurisdiction code."); continue
        if any(m.code == code for m in meas):
            err(where, f"Exchanges based in {code} are described more than once."); continue
        try:
            ws_ = _num(r.get("wash_share")) or 0.0; tv = _num(r.get("turnover"))
        except (TypeError, ValueError):
            err(where, "Wash-trading share and turnover must be numbers."); continue
        tv = 1.0 if tv is None else tv
        if not 0 <= ws_ < 1:
            err(where, "Wash-trading share must be at least 0 and below 1."); continue
        if tv <= 0:
            err(where, "Turnover must be positive."); continue
        meas.append(Measurement(code, ws_, tv, str(r.get("source") or ""), _yes(r.get("verified", ""))))
    # estimation settings
    params = []
    for n, r in rows["Estimation"]:
        where = f"Estimation, row {n}"
        name = str(r.get("name", "")).strip()
        if name not in EST_PARAMS:
            err(where, f"Unknown parameter '{name or 'blank'}'."); continue
        try:
            vals = {k: _num(r.get(k)) for k in ("value", "lower", "upper", "judgement_low", "judgement_high")}
        except (TypeError, ValueError):
            err(where, "Values and bounds must be numbers."); continue
        free = _yes(r.get("free", ""))
        lo = vals["lower"] if vals["lower"] is not None else EST_PARAMS[name]["lo"]
        hi = vals["upper"] if vals["upper"] is not None else EST_PARAMS[name]["hi"]
        if free and not lo < hi:
            err(where, "The lower bound must be below the upper bound."); continue
        if free and vals["value"] is not None and not lo <= vals["value"] <= hi:
            err(where, "The start value must lie within the bounds."); continue
        ps = ParamSpec(name, free, vals["value"], vals["lower"], vals["upper"], vals["judgement_low"], vals["judgement_high"],
                       str(r.get("notes") or ""))
        bad = validate_spec([ps])
        if any(p.name == name for p in params):
            bad.append(f"{name} is listed more than once.")
        for m in bad:
            err(where, m)
        if not bad:
            params.append(ps)
    hyps = []
    for n, r in rows["Hypotheses"]:
        if not str(r.get("id", "")).strip() or not str(r.get("statement", "")).strip():
            err(f"Hypotheses, row {n}", "Each hypothesis needs an ID and a statement."); continue
        if any(h.id == str(r["id"]).strip() for h in hyps):
            err(f"Hypotheses, row {n}", f"ID '{str(r['id']).strip()}' is used more than once."); continue
        hyps.append(Hypothesis(str(r["id"]).strip(), str(r["statement"]), str(r.get("test") or ""), str(r.get("source") or "")))
    # design-level settings from the Settings sheet
    svals = {}
    if "Settings" in wb.sheetnames:
        for row in wb["Settings"].iter_rows(min_row=2, values_only=True):
            if row and len(row) >= 4 and row[3]:
                svals[str(row[3]).strip()] = row[1]
    d = Design(list(eps.values()), obs, meas, params, hyps)
    try:
        if svals.get("estimation_concept") not in (None, ""):
            c = str(svals["estimation_concept"]).strip().lower().strip("()")
            if c not in ("a", "b"):
                raise ValueError
            d.concept = c
        if svals.get("estimation_starts") not in (None, ""):
            d.starts = max(1, min(20, int(float(svals["estimation_starts"]))))
        if svals.get("estimation_seed") not in (None, ""):
            d.seed = int(float(svals["estimation_seed"]))
    except (TypeError, ValueError):
        err("Settings", "Estimation settings: concept must be a or b; starts and seed must be whole numbers.")
    d.label = str(svals.get("design_label") or "")
    d.synthetic = _yes(svals.get("synthetic", ""))
    if obs and not any(o.role == "target" for o in obs):
        warn("Data", "There are no target observations, so nothing can be estimated.")
    if any(i["level"] == "error" for i in issues):
        return None, issues
    return d, issues


def read_workbook(data: bytes):
    """Read the coded world and, if present, the empirical design. Returns (world, design, issues)."""
    world, issues = read_template(data)
    if world is None:
        return None, None, issues
    wb = load_workbook(io.BytesIO(data), data_only=True)
    design, di = read_design(wb, world)
    issues = issues + di
    if any(i["level"] == "error" for i in di):
        return None, None, issues
    return world, design, issues
